import os
import json
import csv
import tempfile
import openpyxl
import pytest
from lab.replay_integrity import verify_exported_artifacts, run_replay_integrity_checks

def test_tampered_tp_in_csv_fails_verifier():
    # PHẦN 08, 106: Mutating take_profit in trades.csv must FAIL verify_exported_artifacts
    with tempfile.TemporaryDirectory() as tmp_dir:
        with open(os.path.join(tmp_dir, "report.json"), "w") as f:
            json.dump({"id": "run-100", "run_id": "run-100", "total_net_pnl": 10.0, "closed_count": 1, "summary": {"realized_net_pnl": 10.0, "closed_trades_count": 1}}, f)

        with open(os.path.join(tmp_dir, "trades.csv"), "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["id", "direction", "entry_price", "stop_loss", "take_profit", "exit_price", "net_pnl", "status"])
            writer.writerow(["t-100", "SHORT", "4123.0", "4133.0", "4151.0", "4103.0", "10.0", "CLOSED"])

        wb = openpyxl.Workbook()
        ws_ov = wb.active
        ws_ov.title = "01_Tong_quan"
        ws_ov.append(["Mã kiểm thử (Run ID)", "run-100"])
        ws_ov.append(["Tổng lợi nhuận thực hiện (Realized Net PnL)", 10.0])

        ws_t = wb.create_sheet(title="03_Chi_tiet_lenh")
        ws_t.append(["Mã lệnh (Trade ID)", "H2", "H3", "H4", "H5", "H6", "H7", "H8", "H9", "Hướng", "H11", "H12", "Entry", "SL", "TP", "H16", "H17", "H18", "H19", "H20", "H21", "H22", "H23", "H24", "H25", "H26", "Lãi ròng ($)", "Realized R"])
        ws_t.append(["t-100", "", "", "", "", "", "", "", "", "SHORT", "", "", 4123.0, 4133.0, 4103.0, "", "", "", "", "", "", "", "", "", "", "", 10.0, 1.0])
        wb.save(os.path.join(tmp_dir, "v12_replay_test.xlsx"))

        result = {
            "id": "run-100",
            "total_net_pnl": 10.0,
            "closed_count": 1,
            "trades": [
                {
                    "id": "t-100",
                    "direction": "SHORT",
                    "entry_price": 4123.0,
                    "stop_loss": 4133.0,
                    "take_profit": 4103.0,
                    "exit_price": 4103.0,
                    "net_pnl": 10.0,
                    "status": "CLOSED"
                }
            ]
        }

        is_valid, errors = verify_exported_artifacts(result, tmp_dir)
        assert is_valid is False
        assert any("TAKE_PROFIT_MISMATCH" in err for err in errors)

def test_nan_in_csv_entry_fails_verifier():
    # PHẦN 08, 106, 139: CSV with NaN entry price must FAIL verifier
    with tempfile.TemporaryDirectory() as tmp_dir:
        with open(os.path.join(tmp_dir, "report.json"), "w") as f:
            json.dump({"id": "run-101", "run_id": "run-101", "total_net_pnl": 10.0, "closed_count": 1, "summary": {"realized_net_pnl": 10.0, "closed_trades_count": 1}}, f)

        with open(os.path.join(tmp_dir, "trades.csv"), "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["id", "direction", "entry_price", "stop_loss", "take_profit", "exit_price", "net_pnl", "status"])
            writer.writerow(["t-101", "LONG", "NaN", "95.0", "110.0", "110.0", "10.0", "CLOSED"])

        result = {
            "id": "run-101",
            "total_net_pnl": 10.0,
            "closed_count": 1,
            "trades": [{"id": "t-101", "direction": "LONG", "entry_price": 100.0, "stop_loss": 95.0, "take_profit": 110.0, "exit_price": 110.0, "net_pnl": 10.0, "status": "CLOSED"}]
        }

        is_valid, errors = verify_exported_artifacts(result, tmp_dir)
        assert is_valid is False
        assert any("NAN" in err or "NON_FINITE" in err or "NON_NUMERIC" in err for err in errors)

def test_unknown_trade_id_in_excel_fails_verifier():
    # PHẦN 08, 106, 139: Fabricated unknown trade ID in Excel must FAIL verifier
    with tempfile.TemporaryDirectory() as tmp_dir:
        with open(os.path.join(tmp_dir, "report.json"), "w") as f:
            json.dump({"id": "run-102", "run_id": "run-102", "total_net_pnl": 10.0, "closed_count": 1, "summary": {"realized_net_pnl": 10.0, "closed_trades_count": 1}}, f)

        with open(os.path.join(tmp_dir, "trades.csv"), "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["id", "direction", "entry_price", "stop_loss", "take_profit", "exit_price", "net_pnl", "status"])
            writer.writerow(["t-102", "LONG", "100.0", "95.0", "110.0", "110.0", "10.0", "CLOSED"])

        wb = openpyxl.Workbook()
        ws_ov = wb.active
        ws_ov.title = "01_Tong_quan"
        ws_ov.append(["Mã kiểm thử (Run ID)", "run-102"])
        ws_ov.append(["Tổng lợi nhuận thực hiện (Realized Net PnL)", 10.0])

        ws_t = wb.create_sheet(title="03_Chi_tiet_lenh")
        ws_t.append(["Mã lệnh (Trade ID)", "H2", "H3", "H4", "H5", "H6", "H7", "H8", "H9", "Hướng", "H11", "H12", "Entry", "SL", "TP", "H16", "H17", "H18", "H19", "H20", "H21", "H22", "H23", "H24", "H25", "H26", "Lãi ròng ($)", "Realized R"])
        # Add an UNKNOWN trade ID "t-FABRICATED"
        ws_t.append(["t-FABRICATED", "", "", "", "", "", "", "", "", "LONG", "", "", 100.0, 95.0, 110.0, "", "", "", "", "", "", "", "", "", "", "", 99999.0, 100.0])
        wb.save(os.path.join(tmp_dir, "v12_replay_test.xlsx"))

        result = {
            "id": "run-102",
            "total_net_pnl": 10.0,
            "closed_count": 1,
            "trades": [{"id": "t-102", "direction": "LONG", "entry_price": 100.0, "stop_loss": 95.0, "take_profit": 110.0, "exit_price": 110.0, "net_pnl": 10.0, "status": "CLOSED"}]
        }

        is_valid, errors = verify_exported_artifacts(result, tmp_dir)
        assert is_valid is False
        assert any("UNKNOWN" in err or "ID_SET_MISMATCH" in err for err in errors)

def test_fake_hash_fails_integrity():
    # PHẦN 08, 62, 105: FAKE hash must FAIL integrity check
    res_dict = {"status": "COMPLETED"}
    closed_trades = [{"id": "t-1", "direction": "LONG", "decision_time": 1000, "entry_time": 2000, "exit_time": 3000, "entry_price": 100.0, "stop_loss": 95.0, "take_profit": 110.0, "status": "CLOSED"}]
    postings = [{"amount": 0.0}]
    initial_eq = 1000.0
    cash = 1000.0

    integrity = run_replay_integrity_checks(
        result_dict=res_dict,
        closed_trades=closed_trades,
        ledger_postings=postings,
        daily_stats={},
        initial_equity=initial_eq,
        reported_cash=cash,
        dataset_hash="FAKEHASH12345678",
        run_config_hash="REALCONFIG12345678"
    )
    assert integrity["status"] == "FAIL"
    ds_check = next(c for c in integrity["checks"] if c["id"] == "dataset_hash_verified")
    assert ds_check["status"] == "FAIL"
