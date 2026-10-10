"""
AURUM DESK — V13.6 NEGATIVE INTEGRITY & ARTIFACTS VERIFICATION TESTS (Phần 07, 87, 88)
Verifies:
- Missing decision_time MUST strictly FAIL integrity checks.
- UNKNOWN direction MUST strictly FAIL pricing geometry checks.
- Tampered CSV trade direction / Net PnL MUST strictly FAIL artifact verifier.
- Tampered XLSX trade cell MUST strictly FAIL artifact verifier.
"""

import os
import json
import csv
import tempfile
import openpyxl
import pytest
from lab.replay_integrity import run_replay_integrity_checks, verify_exported_artifacts


def test_01_fail_on_missing_decision_time():
    # Trade without decision_time (None or missing)
    closed_trades = [
        {"id": "t-1", "decision_time": None, "entry_time": 1050, "exit_time": 2000, "direction": "LONG", "entry_price": 2650.0, "stop_loss": 2645.0, "take_profit": 2665.0}
    ]
    summary = run_replay_integrity_checks(
        result_dict={"closed_count": 1, "wins": 1, "losses": 0, "breakevens": 0},
        closed_trades=closed_trades,
        ledger_postings=[],
        daily_stats={"2026-07-20": {"total_fills": 1}},
        initial_equity=1000.0,
        reported_cash=1000.0,
        dataset_hash="hash-dataset-12345",
        run_config_hash="hash-config-67890"
    )
    assert summary["status"] == "FAIL"
    assert summary["causal_data_ok"] is False
    causal_check = next(c for c in summary["checks"] if c["id"] == "causal_data_order_ok")
    assert causal_check["status"] == "FAIL"
    assert "MISSING_DECISION_TIME" in causal_check["observed"]


def test_02_fail_on_unknown_direction():
    # Trade with UNKNOWN direction
    closed_trades = [
        {"id": "t-1", "decision_time": 1000, "entry_time": 1050, "exit_time": 2000, "direction": "UNKNOWN", "entry_price": 2650.0, "stop_loss": 2645.0, "take_profit": 2665.0}
    ]
    summary = run_replay_integrity_checks(
        result_dict={"closed_count": 1, "wins": 1, "losses": 0, "breakevens": 0},
        closed_trades=closed_trades,
        ledger_postings=[],
        daily_stats={"2026-07-20": {"total_fills": 1}},
        initial_equity=1000.0,
        reported_cash=1000.0,
        dataset_hash="hash-dataset-12345",
        run_config_hash="hash-config-67890"
    )
    assert summary["status"] == "FAIL"
    geom_check = next(c for c in summary["checks"] if c["id"] == "pricing_geometry_verified")
    assert geom_check["status"] == "FAIL"
    assert "INVALID_DIRECTION" in geom_check["observed"]


def test_03_verifier_fails_on_tampered_csv_pnl():
    with tempfile.TemporaryDirectory() as tmp_dir:
        # Create report.json
        with open(os.path.join(tmp_dir, "report.json"), "w") as f:
            json.dump({"total_net_pnl": 10.0, "closed_count": 1, "summary": {"realized_net_pnl": 10.0, "closed_trades_count": 1}}, f)

        # Create trades.csv with tampered PnL 999.0 instead of 10.0
        with open(os.path.join(tmp_dir, "trades.csv"), "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["id", "direction", "entry_price", "stop_loss", "take_profit", "exit_price", "net_pnl", "status"])
            writer.writerow(["t-1", "LONG", "2650.0", "2645.0", "2665.0", "2665.0", "999.0", "CLOSED"])

        # Create minimal xlsx
        wb = openpyxl.Workbook()
        wb.active.title = "01_Tong_quan"
        ws_t = wb.create_sheet(title="03_Chi_tiet_lenh")
        ws_t.append(["H1", "H2", "H3", "H4", "H5", "H6", "H7", "H8", "H9", "Hướng", "H11", "H12", "Entry", "SL", "TP", "H16", "H17", "H18", "H19", "H20", "H21", "H22", "H23", "H24", "H25", "H26", "Lãi ròng ($)"])
        ws_t.append(["t-1", "", "", "", "", "", "", "", "", "LONG", "", "", 2650.0, 2645.0, 2665.0, "", "", "", "", "", "", "", "", "", "", "", 10.0])
        wb.save(os.path.join(tmp_dir, "v12_replay_test.xlsx"))

        result = {
            "total_net_pnl": 10.0,
            "closed_count": 1,
            "trades": [{"id": "t-1", "direction": "LONG", "net_pnl": 10.0, "status": "CLOSED"}]
        }

        is_valid, errors = verify_exported_artifacts(result, tmp_dir)
        assert is_valid is False
        assert any("TRADES_CSV_ROW_0_PNL_MISMATCH" in err or "TRADES_CSV_TOTAL_PNL_MISMATCH" in err for err in errors)


def test_04_verifier_fails_on_tampered_xlsx_direction():
    with tempfile.TemporaryDirectory() as tmp_dir:
        # Create report.json
        with open(os.path.join(tmp_dir, "report.json"), "w") as f:
            json.dump({"total_net_pnl": 10.0, "closed_count": 1, "summary": {"realized_net_pnl": 10.0, "closed_trades_count": 1}}, f)

        # Create valid trades.csv
        with open(os.path.join(tmp_dir, "trades.csv"), "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["id", "direction", "entry_price", "stop_loss", "take_profit", "exit_price", "net_pnl", "status"])
            writer.writerow(["t-1", "LONG", "2650.0", "2645.0", "2665.0", "2665.0", "10.0", "CLOSED"])

        # Create xlsx with tampered Direction "SHORT" instead of "LONG"
        wb = openpyxl.Workbook()
        wb.active.title = "01_Tong_quan"
        ws_t = wb.create_sheet(title="03_Chi_tiet_lenh")
        ws_t.append(["H1", "H2", "H3", "H4", "H5", "H6", "H7", "H8", "H9", "Hướng", "H11", "H12", "Entry", "SL", "TP", "H16", "H17", "H18", "H19", "H20", "H21", "H22", "H23", "H24", "H25", "H26", "Lãi ròng ($)"])
        # Column 10 is "SHORT", but expected is "LONG"
        ws_t.append(["t-1", "", "", "", "", "", "", "", "", "SHORT", "", "", 2650.0, 2645.0, 2665.0, "", "", "", "", "", "", "", "", "", "", "", 10.0])
        wb.save(os.path.join(tmp_dir, "v12_replay_test.xlsx"))

        result = {
            "total_net_pnl": 10.0,
            "closed_count": 1,
            "trades": [{"id": "t-1", "direction": "LONG", "net_pnl": 10.0, "status": "CLOSED"}]
        }

        is_valid, errors = verify_exported_artifacts(result, tmp_dir)
        assert is_valid is False
        assert any("XLSX_ROW_2_DIR_MISMATCH" in err for err in errors)
