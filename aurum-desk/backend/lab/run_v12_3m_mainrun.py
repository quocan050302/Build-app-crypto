"""
Aurum Desk V12: Production Main 3-Month Historical Replay Execution Runner.
- Evaluates real Bitget Classic Futures XAUUSDT data from 2026-07-09 22:00:00 to 2026-10-09 22:00:00 UTC+7.
- Complete 50-day warmup history (from 2026-05-20 22:00:00 UTC+7).
- Initial Equity: 1,000 USDT | Leverage: 30x ISOLATED | Risk: 0.25% equity.
- Exports comprehensive 10-sheet Excel workbook (.xlsx) and JSON/HTML/CSV machine artifacts.
"""
import os
import sys
import json
import time
from datetime import datetime
from zoneinfo import ZoneInfo

# Add backend directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import schemas
from lab.replay_engine import ReplayEngine
from lab.historical_market_data import (
    HistoricalMarketDataProvider,
    subtract_calendar_months,
    compute_dataset_hash,
    VN_TZ
)
import openpyxl

def run_main_3m_v12():
    print("=" * 80)
    print("AURUM DESK V12 — 3-MONTH REAL HISTORICAL REPLAY EXECUTION")
    print("=" * 80)

    # 1. Define exact timestamps
    cutoff_dt = datetime(2026, 10, 9, 22, 0, 0, tzinfo=VN_TZ)
    start_dt = subtract_calendar_months(cutoff_dt, 3)  # Exactly 2026-07-09 22:00:00
    cutoff_ms = int(cutoff_dt.timestamp() * 1000)
    start_ms = int(start_dt.timestamp() * 1000)
    warmup_days = 50
    warmup_ms = start_ms - (warmup_days * 24 * 3600 * 1000)
    warmup_dt = datetime.fromtimestamp(warmup_ms / 1000.0, tz=VN_TZ)

    print(f"Cutoff Time  (End)  : {cutoff_dt.strftime('%Y-%m-%d %H:%M:%S %Z')} ({cutoff_ms})")
    print(f"Start Time   (Start): {start_dt.strftime('%Y-%m-%d %H:%M:%S %Z')} ({start_ms})")
    print(f"Warmup Start (50D)  : {warmup_dt.strftime('%Y-%m-%d %H:%M:%S %Z')} ({warmup_ms})")
    print(f"Period Days Difference: {(cutoff_dt - start_dt).days} days (Exact 3 calendar months)")

    # 2. Prepare Replay Request
    req = schemas.ReplayRunRequest(
        run_name="Aurum_V12_Main_3Months_XAUUSDT",
        symbol="XAUUSDT",
        mode="HISTORICAL_MARKET",
        start_ts=start_ms,
        end_ts=cutoff_ms,
        warmup_days=warmup_days,
        initial_equity=1000.0,
        risk_pct=0.25,
        leverage=30,
        margin_mode="ISOLATED",
        fee_rate=0.0006,
        spread_multiplier=1.0,
        slippage_multiplier=1.0,
        export_artifacts=True
    )

    t0 = time.time()
    print("\nExecuting replay engine over verified Bitget historical data...")
    res = ReplayEngine.run_replay(req)
    t_elapsed = time.time() - t0

    print(f"Execution completed in {t_elapsed:.2f} seconds.")
    print("-" * 80)
    print("REPLAY SUMMARY METRICS:")
    print(f"  Run ID                  : {res.id}")
    print(f"  Dataset Type            : {res.dataset_type}")
    print(f"  Execution Fidelity      : {res.execution_fidelity}")
    print(f"  Initial Equity          : ${res.initial_equity:.2f} USDT")
    print(f"  Final Equity            : ${res.final_equity:.2f} USDT")
    print(f"  Cash Balance            : ${getattr(res, 'cash_balance', res.final_equity):.2f} USDT")
    print(f"  Open MTM                : ${getattr(res, 'open_mtm', 0.0):.2f} USDT")
    print(f"  Realized Net PnL        : ${res.total_net_pnl:.2f} USDT")
    print(f"  Total Trades (Fills)    : {res.total_trades}")
    print(f"  Wins / Losses / Evens   : {res.wins}W / {res.losses}L / {res.breakevens}BE")
    print(f"  Win Rate (Closed)       : {res.win_rate_pct:.1f}%")
    pf_str = f"{res.profit_factor:.2f}" if res.profit_factor is not None else "N/A (0 Losses)"
    print(f"  Profit Factor           : {pf_str}")
    print(f"  Expectancy R            : {res.expectancy_r:.2f}R")
    print(f"  Max Drawdown (MTM)      : {res.max_drawdown_pct:.2f}% (${res.max_drawdown_usdt:.2f})")
    print(f"  Max Consecutive Losses  : {res.max_consecutive_losses}")
    print(f"  Total Fees              : ${res.total_fees:.2f} USDT")
    print(f"  Total Slippage          : ${res.total_slippage:.2f} USDT")
    print(f"  Signals (READY) Count   : {res.signals_count}")
    print(f"  Policy Rejected Count   : {res.rejected_count}")
    print(f"  Rejection Reasons       : {res.rejection_reasons}")
    print("-" * 80)

    # 3. Verify Artifacts
    print(f"Artifacts Directory: {res.artifacts_dir}")
    if res.artifacts_dir and os.path.exists(res.artifacts_dir):
        files = sorted(os.listdir(res.artifacts_dir))
        print("Generated Files:")
        excel_path = None
        for f in files:
            size_kb = os.path.getsize(os.path.join(res.artifacts_dir, f)) / 1024.0
            print(f"  - {f} ({size_kb:.1f} KB)")
            if f.endswith(".xlsx"):
                excel_path = os.path.join(res.artifacts_dir, f)

        if excel_path:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            print("\nExcel Workbook Verification:")
            print(f"  File: {os.path.basename(excel_path)}")
            print(f"  Sheets ({len(wb.sheetnames)}): {wb.sheetnames}")
            for sname in wb.sheetnames:
                ws = wb[sname]
                print(f"    * {sname}: {ws.max_row} rows, {ws.max_column} cols")
            wb.close()

    print("=" * 80)
    print("V12 3-MONTH REPLAY RUN FINISHED SUCCESSFULLY!")
    print("=" * 80)
    return res

if __name__ == "__main__":
    run_main_3m_v12()
